# -*- coding: utf-8 -*-
from __future__ import annotations

import random
from collections import Counter, defaultdict
from copy import copy
from decimal import Decimal, ROUND_FLOOR
from pathlib import Path

from openpyxl import Workbook, load_workbook

SRC = Path(r"E:\学习资料\研\我的论文\金融科技\第二章论文\标注数据\advisor_recall_weighted_pool_sample_2014_2026_part1_人工标注.xlsx")
OUT = SRC.with_name(SRC.stem + "_按sample_pool加权抽样1000.xlsx")
TARGET = 1000
SEED = 20260612
POOL_COL = "sample_pool"


def copy_cell(source_cell, target_cell) -> None:
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)
    target_cell.number_format = source_cell.number_format
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.fill = copy(source_cell.fill)
    target_cell.font = copy(source_cell.font)
    target_cell.border = copy(source_cell.border)


def main() -> None:
    wb = load_workbook(SRC)
    ws = wb[wb.sheetnames[0]]
    headers = [cell.value for cell in ws[1]]
    if POOL_COL not in headers:
        raise ValueError(f"missing column: {POOL_COL}")
    pool_idx = headers.index(POOL_COL) + 1

    rows_by_pool: dict[str, list[int]] = defaultdict(list)
    for row_idx in range(2, ws.max_row + 1):
        value = ws.cell(row_idx, pool_idx).value
        if value is None or str(value).strip() == "":
            raise ValueError(f"empty sample_pool at row {row_idx}")
        rows_by_pool[str(value).strip()].append(row_idx)

    total_rows = sum(len(v) for v in rows_by_pool.values())
    if total_rows < TARGET:
        raise ValueError(f"not enough rows: {total_rows} < {TARGET}")

    quotas: dict[str, int] = {}
    remainders: list[tuple[Decimal, str]] = []
    assigned = 0
    for pool in sorted(rows_by_pool):
        exact = Decimal(len(rows_by_pool[pool])) * Decimal(TARGET) / Decimal(total_rows)
        floor_value = int(exact.to_integral_value(rounding=ROUND_FLOOR))
        quotas[pool] = floor_value
        assigned += floor_value
        remainders.append((exact - Decimal(floor_value), pool))

    for _, pool in sorted(remainders, key=lambda x: (-x[0], x[1]))[: TARGET - assigned]:
        quotas[pool] += 1

    rng = random.Random(SEED)
    selected_rows: list[int] = []
    for pool in sorted(rows_by_pool):
        pool_rows = rows_by_pool[pool][:]
        rng.shuffle(pool_rows)
        selected_rows.extend(pool_rows[: quotas[pool]])
    selected_rows.sort()

    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = ws.title

    for target_col, source_cell in enumerate(ws[1], start=1):
        copy_cell(source_cell, out_ws.cell(1, target_col))

    for out_row_idx, source_row_idx in enumerate(selected_rows, start=2):
        for col_idx in range(1, ws.max_column + 1):
            copy_cell(ws.cell(source_row_idx, col_idx), out_ws.cell(out_row_idx, col_idx))

    for col_idx, dimension in ws.column_dimensions.items():
        out_ws.column_dimensions[col_idx].width = dimension.width
        out_ws.column_dimensions[col_idx].hidden = dimension.hidden

    out_ws.freeze_panes = ws.freeze_panes
    if ws.auto_filter and ws.auto_filter.ref:
        out_ws.auto_filter.ref = f"A1:{out_ws.cell(out_ws.max_row, out_ws.max_column).coordinate}"

    out_wb.save(OUT)

    check_wb = load_workbook(OUT, read_only=True, data_only=False)
    check_ws = check_wb[check_wb.sheetnames[0]]
    check_headers = [cell.value for cell in next(check_ws.iter_rows(min_row=1, max_row=1))]
    check_pool_idx = check_headers.index(POOL_COL)
    check_counts = Counter(
        str(row[check_pool_idx]).strip()
        for row in check_ws.iter_rows(min_row=2, values_only=True)
    )
    print(f"SOURCE={SRC}")
    print(f"OUTPUT={OUT}")
    print(f"SEED={SEED}")
    print(f"SOURCE_ROWS={total_rows}")
    print(f"OUTPUT_ROWS={check_ws.max_row - 1}")
    print("QUOTAS=" + ",".join(f"{pool}:{quotas[pool]}" for pool in sorted(quotas)))
    print("OUTPUT_COUNTS=" + ",".join(f"{pool}:{check_counts[pool]}" for pool in sorted(check_counts)))
    print("VALIDATION=" + ("PASS" if check_ws.max_row - 1 == TARGET and dict(check_counts) == quotas else "FAIL"))


if __name__ == "__main__":
    main()
